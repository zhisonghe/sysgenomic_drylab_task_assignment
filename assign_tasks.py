#!/usr/bin/env python3
"""
Reproducible assignment of drylab tasks to student groups.

Algorithm (rank-by-rank lottery with tiers):
  0. Parse each group's preference list into tiers. Consecutive tasks from the
     same paper form one tier (equally preferred). Parenthetical hints such as
     "(medium / neuroscience)" are kept as fallback metadata.
  1. Explicit preferences. For round r = 1, 2, ...: every unassigned group
     applies to the still-free tasks in its r-th tier. Within the round, the
     largest possible number of applicants is served; who is served is decided
     by a fresh seeded lottery. List length gives no priority: everyone in
     round r competes equally. Within its tier, a served group gets the task
     that the fewest other unassigned groups list (lottery among ties), so an
     indifferent choice never needlessly takes another group's later choice.
  2. Fallback hints. Groups whose lists are exhausted but who gave hints rank
     the remaining free tasks by how well they match (difficulty distance +
     topic mismatch), then go through the same round-by-round lottery.
  3. Everyone else (exhausted lists without hints, and groups with no
     preference) receives the leftover tasks at random.

Usage: python3 assign_tasks.py [--simulate N]
"""
import argparse
import random
import re
from collections import Counter, defaultdict

import openpyxl

SEED = 1012705090  # drawn once with secrets.randbelow(2**31), then fixed
XLSX = "Systems Genomics Organisation 2026.xlsx"
OUT = "Task assignment 2026.xlsx"

DIFFICULTY = {"low": 1, "low-medium": 2, "medium": 3, "medium-high": 4, "high": 5}

# Tasks counted as "neuroscience" for the topic hint: all brain / nervous-system
# datasets. 5 (mouse cortical layer 4) and 6 (bipolar disorder, post-mortem
# brain) are included although their task text does not say "brain".
NEURO_TASKS = {4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 18, 19, 20, 21, 22, 23, 25, 26}
TOPICS = {"neuroscience": NEURO_TASKS}


# ---------------------------------------------------------------- parsing ----

def load(path=XLSX):
    wb = openpyxl.load_workbook(path, data_only=True)

    tasks = {}
    ws = wb["Datasets"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], int):
            continue
        level = row[7].split("(")[0].strip().lower()
        tasks[row[0]] = {"paper": row[1].strip(), "task": row[3].strip(),
                         "complexity": row[7], "level": DIFFICULTY[level]}

    groups = {}
    ws = wb["Drylab Groups"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not isinstance(row[0], int):
            continue
        groups[row[0]] = parse_preference(row[12], tasks)
    return tasks, groups


def parse_preference(cell, tasks):
    """Return {'raw', 'tiers': [[task, ...], ...], 'levels': set, 'topics': set}."""
    pref = {"raw": "" if cell is None else str(cell), "tiers": [], "levels": set(), "topics": set()}
    if cell is None:
        return pref
    text = str(cell)
    hint = re.search(r"\((.*)\)", text)
    if hint:
        for token in re.split(r"\s*(?:/|\bor\b|,)\s*", hint.group(1).lower()):
            if token in DIFFICULTY:
                pref["levels"].add(DIFFICULTY[token])
            elif token in TOPICS:
                pref["topics"].add(token)
            elif token:
                raise ValueError(f"Unknown hint {token!r} in {text!r}")
        text = text[:hint.start()]
    ids = [int(x) for x in re.findall(r"\d+", text)]
    for t in ids:
        if pref["tiers"] and tasks[pref["tiers"][-1][-1]]["paper"] == tasks[t]["paper"]:
            pref["tiers"][-1].append(t)
        else:
            pref["tiers"].append([t])
    return pref


def fallback_tiers(pref, free, tasks):
    """Rank free tasks by match to the hints; equal scores form one tier."""
    def score(t):
        s = min(abs(tasks[t]["level"] - l) for l in pref["levels"]) if pref["levels"] else 0
        s += sum(t not in TOPICS[topic] for topic in pref["topics"])
        return s
    by_score = defaultdict(list)
    for t in sorted(free):
        by_score[score(t)].append(t)
    return [by_score[s] for s in sorted(by_score)]


# --------------------------------------------------------------- matching ----

def can_match(groups, options, taken=frozenset()):
    """True if every group in `groups` can get a distinct task from options[g]."""
    owner = {}

    def augment(g, seen):
        for t in options[g]:
            if t in taken or t in seen:
                continue
            seen.add(t)
            if t not in owner or augment(owner[t], seen):
                owner[t] = g
                return True
        return False

    return all(augment(g, set()) for g in groups)


def lottery_round(options, rng, wanted_by):
    """
    options: {group: [free tasks in its current tier]}.
    wanted_by: {task: set of unassigned groups listing it anywhere}.
    Serve as many groups as possible. A random priority order decides who is
    served when not everyone can be. Each served group then gets, among the
    tasks in its tier that keep the other served groups feasible, the one
    wanted by the fewest groups not served in this round (random among ties).
    Returns ({group: task}, priority_order).
    """
    order = sorted(options)
    rng.shuffle(order)
    served = []
    for g in order:
        if can_match(served + [g], options):
            served.append(g)
    result, taken = {}, set()
    for g in served:
        rest = [h for h in served if h not in result and h != g]
        choices = list(options[g])
        rng.shuffle(choices)
        feasible = [t for t in choices
                    if t not in taken and can_match(rest, options, frozenset(taken | {t}))]
        # min() keeps the first of equal values, i.e. the shuffled order breaks ties
        t = min(feasible, key=lambda t: len(wanted_by.get(t, set()) - set(served)))
        result[g] = t
        taken.add(t)
    return result, order


def run_rounds(tier_lists, free, assignment, outcome, label, rng, log, explicit):
    """
    Round-by-round lottery over tier_lists ({group: [tier, ...]}). `explicit`
    holds every group's stated tiers, used to see which tasks unassigned groups
    still want.
    """
    pending = set(tier_lists)
    for r in range(max((len(v) for v in tier_lists.values()), default=0)):
        options = {}
        for g in sorted(pending):
            if r < len(tier_lists[g]):
                opts = [t for t in tier_lists[g][r] if t in free]
                if opts:
                    options[g] = opts
        if not options:
            continue
        wanted_by = defaultdict(set)
        for g, tiers in explicit.items():
            if g not in assignment:
                for t in (t for tier in tiers for t in tier):
                    wanted_by[t].add(g)
        won, order = lottery_round(options, rng, wanted_by)
        demand = Counter(t for opts in options.values() for t in opts)
        for g, opts in options.items():
            contested = len(options) > 1 and any(demand[t] > 1 for t in opts)
            log.append({"stage": label, "round": r + 1, "group": g,
                        "applied_to": ", ".join(map(str, opts)),
                        "lottery_position": order.index(g) + 1 if contested else "",
                        "result": won.get(g, "lost")})
        for g, t in won.items():
            assignment[g] = t
            outcome[g] = f"{label} {r + 1}"
            free.discard(t)
            pending.discard(g)


def assign(tasks, groups, seed=SEED):
    rng = random.Random(seed)
    free, assignment, outcome, log = set(tasks), {}, {}, []

    # Stage 1: explicit preferences
    explicit = {g: p["tiers"] for g, p in groups.items() if p["tiers"]}
    run_rounds(explicit, free, assignment, outcome, "Preference tier", rng, log, explicit)

    # Stage 2: fallback hints for groups whose lists are exhausted
    hinted = {g: fallback_tiers(p, free, tasks) for g, p in groups.items()
              if g not in assignment and (p["levels"] or p["topics"])}
    run_rounds(hinted, free, assignment, outcome, "Fallback tier", rng, log, explicit)

    # Stage 3: random leftovers
    rest = sorted(g for g in groups if g not in assignment)
    left = sorted(free)
    rng.shuffle(left)
    for g, t in zip(rest, left):
        assignment[g] = t
        outcome[g] = "Random (no preference)" if not groups[g]["tiers"] else "Random (list exhausted)"
        log.append({"stage": "Random", "round": "", "group": g, "applied_to": "any",
                    "lottery_position": "", "result": t})

    assert sorted(assignment.values()) == sorted(tasks) and len(assignment) == len(groups)
    return assignment, outcome, log


# ----------------------------------------------------------------- output ----

def simulate(tasks, groups, n):
    counts = {g: Counter() for g in groups}
    for s in range(n):
        _, outcome, _ = assign(tasks, groups, seed=s)
        for g, o in outcome.items():
            counts[g][o] += 1
    return counts


def write_output(tasks, groups, assignment, outcome, log, sim, n_sim):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assignment"
    ws.append(["Group #", "Preference (raw)", "Parsed tiers", "Hints", "Assigned task",
               "Paper", "Task", "Task complexity", "Outcome"])
    for g in sorted(groups):
        p, t = groups[g], assignment[g]
        hints = sorted(k for k, v in DIFFICULTY.items() if v in p["levels"]) + sorted(p["topics"])
        ws.append([g, p["raw"], " > ".join("{" + ", ".join(map(str, tier)) + "}" for tier in p["tiers"]),
                   ", ".join(hints), t, tasks[t]["paper"], tasks[t]["task"],
                   tasks[t]["complexity"], outcome[g]])

    ws = wb.create_sheet("Lottery log")
    ws.append([f"Seed: {SEED}"])
    ws.append(["Stage", "Round", "Group #", "Applied to (free tasks in tier)",
               "Lottery position (contested only)", "Result"])
    for e in log:
        ws.append([e["stage"], e["round"], e["group"], e["applied_to"], e["lottery_position"], e["result"]])

    if sim:
        cats = sorted({o for c in sim.values() for o in c},
                      key=lambda o: (o.startswith("Random"), o.startswith("Fallback"),
                                     int(o.split()[-1]) if o[-1].isdigit() else 0, o))
        ws = wb.create_sheet("Fairness check")
        ws.append([f"Share of {n_sim} runs with seeds 0..{n_sim - 1} (official run uses seed {SEED})"])
        ws.append(["Group #", "Preference (raw)"] + cats)
        for g in sorted(groups):
            ws.append([g, groups[g]["raw"]] + [round(sim[g][c] / n_sim, 3) for c in cats])

    for sheet in wb:
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.save(OUT)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--simulate", type=int, default=10000,
                    help="number of seeds for the fairness check (0 to skip)")
    args = ap.parse_args()

    tasks, groups = load()
    assignment, outcome, log = assign(tasks, groups)
    sim = simulate(tasks, groups, args.simulate) if args.simulate else None
    write_output(tasks, groups, assignment, outcome, log, sim, args.simulate)

    for g in sorted(groups):
        print(f"Group {g:>2} -> task {assignment[g]:>2}  [{outcome[g]}]  pref: {groups[g]['raw']}")
    print(f"\nWritten to {OUT!r}")


if __name__ == "__main__":
    main()
